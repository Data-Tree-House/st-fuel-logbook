import datetime
from pathlib import Path
from typing import get_args

import pytest
from openpyxl import load_workbook
from sqlalchemy.engine import Engine

from db import crud
from db.model import FuelTypeLiteral
from utils.bulk.template import generate_bulk_template
from utils.bulk.types import SheetNames, all_sheet_names


class TestNewCar:
    @pytest.fixture(autouse=True)
    def setup(self, mock_engine: Engine):
        self.user_id = "test_sub"
        crud.upsert_user(
            sub=self.user_id,
            name="Test User",
            email="johan@datatreehouse.org",
            picture="https://example.com/profile.jpg",
            engine=mock_engine,
        )
        self.car_id = crud.new_car(
            user_id=self.user_id,
            nickname="Ford Focus",
            fuel_type="Unleaded Petrol 95",
            engine=mock_engine,
        )
        crud.new_fuel_entry(
            self.user_id,
            mock_engine,
            car_id=self.car_id,
            entry_datetime=datetime.datetime(2026, 6, 1, 12, 0, tzinfo=datetime.UTC),
            odometer=10000.0,
            trip=500.0,
            fuel_filled=50.0,
            price=750.0,
            currency="ZAR",
            location="Cape Town",
        )

    def test_generate_bulk_template(
        self,
        mock_engine: Engine,
        temp_dir: Path,
    ):
        tmp_path = temp_dir / "bulk_template.xlsx"
        all_cars = crud.get_all_cars(self.user_id, mock_engine)
        all_fuel_entries = crud.get_all_fuel_entries(self.user_id, mock_engine)

        assert all_cars is not None
        assert all_fuel_entries is not None

        wb = generate_bulk_template(
            fuel_entries=all_fuel_entries,
            cars=all_cars,
            dir=temp_dir,
        )
        wb.save(tmp_path)

        assert tmp_path.exists()

        # ~.~.~.~.~.~.~.~ // CHECK SHEETS // ~.~.~.~.~.~.~.~

        wb = load_workbook(tmp_path)

        for required_sheet in all_sheet_names:
            assert required_sheet in wb.sheetnames

        # ~.~.~.~.~.~.~.~ // CHECK METADATA VALUES // ~.~.~.~.~.~.~.~

        metadata_sheet = wb[SheetNames.METADATA_SHEET]
        required_values = list(get_args(FuelTypeLiteral))

        for i, expected_value in enumerate(required_values, start=1):
            cell_value = metadata_sheet[f"A{i}"].value
            assert cell_value == expected_value, f"Expected '{expected_value}' in A{i}, got '{cell_value}'"
