import tempfile
from collections.abc import Sequence
from pathlib import Path
from typing import get_args

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import quote_sheetname
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from db.model import Base, Car, FuelEntry, FuelTypeLiteral

from .types import ROW_ACTION_COLUMN_NAME, DefinedNames, RowAction, SheetNames, all_row_actions, all_sheet_names


def sql_to_excel(record: Base) -> dict:
    return {k: getattr(record, v) for k, v in record.excel_names_to_columns().items()}


def generate_bulk_template(
    fuel_entries: Sequence[FuelEntry],
    cars: Sequence[Car],
    dir: Path | None = None,
) -> Workbook:
    with tempfile.TemporaryDirectory(
        prefix="bulk-template",
        dir=dir,
    ) as tmpdir:
        tmp_path: Path = Path(tmpdir) / "bulk_template.xlsx"

        wb = Workbook()

        # ====> (1) CREATE ALL THE SHEETS
        _create_sheets(wb)

        # ====> (2) ADD FUEL TYPE AND ROW ACTION DROPDOWNS
        fuel_type_dv = define_dropdown_options(
            wb=wb,
            options=list(get_args(FuelTypeLiteral)),
            sheet_name=SheetNames.METADATA_SHEET,
            column="A",
            defined_name=DefinedNames.FUEL_TYPE_OPTIONS,
            start_row=1,
            allow_more_options=True,
        )
        row_action_dv = define_dropdown_options(
            wb=wb,
            options=list(all_row_actions),
            sheet_name=SheetNames.METADATA_SHEET,
            column="B",
            defined_name=DefinedNames.ROW_ACTION_OPTIONS,
            start_row=1,
        )

        # ====> (3) SAVE WHAT WE GOT SO FAR
        # At this point, it will be easier to use pandas to save and manipulate the fuel entries
        wb.save(tmp_path)

        # ====> (4) DUMP FUEL ENTRIES
        dump_fuel_entries(
            fuel_entries=fuel_entries,
            tmp_path=tmp_path,
            row_action_dv=row_action_dv,
        )

        return load_workbook(tmp_path)


def _create_sheets(wb: Workbook) -> None:
    ws = wb.active
    if ws is None:
        raise ValueError("No active sheet found")

    ws.title = all_sheet_names[0]

    for new_sheet in all_sheet_names[1:]:
        wb.create_sheet(new_sheet)


def define_dropdown_options(
    wb: Workbook,
    options: list[str],
    sheet_name: str,
    column: str,
    defined_name: str,
    start_row: int = 1,
    allow_more_options: bool = False,
) -> DataValidation:
    """
    Writes dropdown options to a sheet, registers a named range,
    and returns a DataValidation object ready to be applied.
    """
    num_values = len(options)
    ws = wb[sheet_name]

    # Populate values into the sheet
    for i, value in enumerate(options):
        ws[f"{column}{start_row + i}"] = value

    # Register the named range pointing at the written values
    if allow_more_options:
        ref = f"{quote_sheetname(sheet_name)}!${column}:${column}"
    else:
        ref = f"{quote_sheetname(sheet_name)}!${column}${start_row}:${column}${start_row + num_values - 1}"
    wb.defined_names.add(DefinedName(name=defined_name, attr_text=ref))

    # Build and return the DataValidation (not yet applied to any cell)
    return DataValidation(
        type="list",
        formula1=defined_name,
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="Invalid input",
        error="Please select a value from the list.",
    )


def apply_dropdown(
    wb: Workbook,
    dv: DataValidation,
    cell_range: str,
    sheet_name: str | None = None,
) -> None:
    """
    Applies a DataValidation to a cell range on a sheet.
    Defaults to the active sheet if sheet_name is not provided.
    """
    if sheet_name is not None:
        ws = wb[sheet_name]
    else:
        ws = wb.active
        assert ws is not None, "No active sheet found"

    ws.add_data_validation(dv)
    dv.add(cell_range)


def num_to_col(n: int) -> str:
    result = ""
    n += 1  # 0 index offset
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        result = chr(65 + remainder) + result
    return result


def dump_fuel_entries(
    fuel_entries: Sequence[FuelEntry],
    tmp_path: Path,
    row_action_dv: DataValidation,
) -> None:
    # ====> CREATE DATA
    data = [sql_to_excel(r) for r in fuel_entries]
    df = pd.DataFrame(data)

    # ====> ADD A COLUMN ROW_ACTION_COLUMN_NAME AND SET THEM ALL TO "READ"
    df[ROW_ACTION_COLUMN_NAME] = RowAction.READ
    action_col_index = df.columns.get_loc(ROW_ACTION_COLUMN_NAME)
    assert isinstance(action_col_index, int), "Expected an integer column index"
    action_column_letter = num_to_col(int(action_col_index))

    # ====> SAVE
    with pd.ExcelWriter(tmp_path, mode="a", engine="openpyxl", if_sheet_exists="replace") as writer:
        df.to_excel(
            writer,
            sheet_name=SheetNames.FUEL_ENTRIES_SHEET,
            index=False,
        )

    # ====> ADD Column with row actions
    wb = load_workbook(tmp_path)
    fuel_entries_ws = wb[SheetNames.FUEL_ENTRIES_SHEET]
    max_row = fuel_entries_ws.max_row

    apply_dropdown(
        wb=wb,
        dv=row_action_dv,
        cell_range=f"{action_column_letter}2:{action_column_letter}{max_row}",
        sheet_name=SheetNames.FUEL_ENTRIES_SHEET,
    )

    wb.save(tmp_path)
